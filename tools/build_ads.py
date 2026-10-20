# -*- coding: utf-8 -*-
"""Génère la campagne Google Ads Search (CPC manuel) dans google-ads/.

Usage : python3 tools/build_ads.py
Produit des fichiers CSV importables dans Google Ads Editor
(Compte > Importer > Coller du texte / Depuis un fichier) et vérifie
les limites de caractères de Google (titres 30, descriptions 90, chemins 15).
"""
import csv
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from site_data import SITE, COMMUNES  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "google-ads")
BASE = SITE["base_url"]

C1 = "Search - Plombier Toulouse - Urgence"
C2 = "Search - Plombier Communes 25 km"

# ---------------------------------------------------------------------------
# Groupes d'annonces : (campagne, nom, url, enchère exacte, enchère expression, mots-clés, titres, descriptions, chemin1, chemin2)
# Les enchères sont des points de départ (en €), à ajuster après 7 à 14 jours.
# ---------------------------------------------------------------------------
COMMON_DESC = [
    "Plombier disponible 24h/24 et 7j/7, nuit, week-end et jours fériés. Appelez maintenant !",
    "Devis écrit remis sur place avant travaux. Rien ne commence sans votre accord.",
    "Toulouse et 25 km autour : Blagnac, Colomiers, Tournefeuille, Balma, Muret...",
    "Facture détaillée pour votre assurance. Paiement CB, espèces ou chèque.",
]

AD_GROUPS = [
    (C1, "Plombier Urgence Toulouse", "/", 7.50, 6.00,
     ["plombier urgence toulouse", "urgence plombier toulouse", "plombier urgent toulouse", "sos plombier toulouse",
      "plombier 24h 24 toulouse", "plombier 24 24 toulouse", "plombier de garde toulouse", "plombier nuit toulouse",
      "plombier dimanche toulouse", "plombier week end toulouse", "plombier jour ferie toulouse", "depannage plomberie urgence toulouse",
      "plombier urgence 31", "plombier urgence haute garonne"],
     ["Plombier Urgence Toulouse", "SOS Plombier 24h/24 7j/7", "Plombier en 30 min Environ", "Appelez le 09 70 70 71 29",
      "Devis Écrit Avant Travaux", "Plombier Nuit et Week-end", "Intervention Rapide", "Dépannage Plomberie Urgent",
      "Disponible Dimanche et Férié", "Fuite, Bouchon, Chauffe-eau", "Plombiers Qualifiés", "Facture pour l'Assurance",
      "Toulouse et 25 km Autour", "Un Plombier Vous Répond", "Paiement CB Accepté"],
     ["Urgence plomberie à Toulouse ? Un plombier se déplace rapidement, de jour comme de nuit.",
      "Fuite d'eau, WC bouché, chauffe-eau en panne : intervention 24h/24 à Toulouse et autour."] + COMMON_DESC[1:3],
     "Plombier", "Urgence"),

    (C1, "Plombier Toulouse", "/", 6.00, 4.80,
     ["plombier toulouse", "plombier a toulouse", "plombier toulouse 31", "plombier 31000", "plombier 31", "artisan plombier toulouse",
      "entreprise plomberie toulouse", "plomberie toulouse", "depannage plomberie toulouse", "plombier depannage toulouse",
      "plombier toulouse centre", "bon plombier toulouse", "plombier toulouse devis", "plombier sanitaire toulouse"],
     ["Plombier à Toulouse 24h/24", "Dépannage Plomberie Toulouse", "Devis Gratuit Avant Travaux", "Appelez le 09 70 70 71 29",
      "Plombier en 30 min Environ", "Tous Quartiers de Toulouse", "Fuite, Bouchon, Chauffe-eau", "Intervention Rapide",
      "Plombiers Qualifiés", "Toutes Marques", "Facture pour l'Assurance", "Disponible 7j/7",
      "Toulouse et 25 km Autour", "Prix Annoncé Avant Travaux", "Paiement CB Accepté"],
     ["Votre plombier à Toulouse : fuite, débouchage, chauffe-eau. Intervention 24h/24 et 7j/7.",
      "Un plombier qualifié évalue votre panne au téléphone et se déplace avec le bon matériel."] + COMMON_DESC[1:3],
     "Plombier", "Toulouse"),

    (C1, "Plombier Pres de Moi", "/", 5.50, 0,
     ["plombier", "plombier urgence", "plombier pres de moi", "plombier autour de moi", "plombier proche de moi",
      "plombier ouvert maintenant", "sos plombier", "urgence plombier", "depannage plomberie", "plombier 24h 24",
      "plombier de garde", "plombier dimanche", "plombier nuit"],
     ["Plombier Près de Chez Vous", "Plombier Urgence 24h/24", "Intervention Rapide", "Appelez le 09 70 70 71 29",
      "Devis Écrit Avant Travaux", "Plombier Disponible Maintenant", "Nuit, Week-end et Fériés", "Toulouse et 25 km Autour",
      "Fuite, Bouchon, Chauffe-eau", "Plombiers Qualifiés", "Facture pour l'Assurance", "Un Plombier Vous Répond",
      "SOS Plombier Toulouse", "Toutes Marques", "Paiement CB Accepté"],
     ["Besoin d'un plombier maintenant ? Nous intervenons à Toulouse et 25 km autour, 24h/24.",
      COMMON_DESC[0], COMMON_DESC[1], COMMON_DESC[3]],
     "Plombier", "Urgence"),

    (C1, "Fuite d'eau", "/fuite-d-eau-toulouse/", 5.00, 4.00,
     ["fuite d'eau toulouse", "fuite eau toulouse", "plombier fuite d'eau toulouse", "reparation fuite d'eau toulouse",
      "fuite d'eau urgence", "fuite d'eau plombier", "fuite canalisation toulouse", "fuite tuyau toulouse", "fuite sous evier",
      "fuite chasse d'eau", "reparer fuite d'eau", "fuite d'eau appartement", "fuite d'eau plafond", "fuite d'eau maison"],
     ["Fuite d'Eau à Toulouse ?", "Plombier Fuite d'Eau 24h/24", "Fuite Stoppée Rapidement", "Appelez le 09 70 70 71 29",
      "Réparation Dès le 1er Passage", "Devis Écrit Avant Travaux", "Coupez l'Eau et Appelez-Nous", "Plombier en 30 min Environ",
      "Canalisation, Flexible, Joint", "Facture pour l'Assurance", "Nuit, Week-end et Fériés", "Toulouse et 25 km Autour",
      "Intervention Rapide", "Plombiers Qualifiés", "Fuite Sous Évier ou Plafond"],
     ["Fuite d'eau ? Coupez l'eau et appelez : un plombier stoppe la fuite et répare, 24h/24.",
      "Flexible, joint, raccord ou canalisation : réparé dès le premier passage le plus souvent."] + COMMON_DESC[1:3],
     "Fuite-d-eau", "Toulouse"),

    (C1, "Recherche de fuite", "/recherche-de-fuite-toulouse/", 4.50, 3.60,
     ["recherche de fuite toulouse", "recherche fuite toulouse", "detection fuite toulouse", "detection de fuite d'eau toulouse",
      "recherche de fuite non destructive", "recherche fuite non destructive toulouse", "recherche de fuite d'eau", "fuite d'eau cachee",
      "detection fuite canalisation", "recherche fuite plancher chauffant", "camera thermique fuite", "recherche de fuite assurance",
      "fuite d'eau compteur qui tourne", "fuite d'eau introuvable"],
     ["Recherche de Fuite Toulouse", "Détection Sans Casse", "Fuite Cachée ? On la Trouve", "Caméra Thermique, Gaz Traceur",
      "Rapport pour l'Assurance", "Appelez le 09 70 70 71 29", "Recherche Non Destructive", "Devis Écrit Avant Travaux",
      "Fuite Sous Dalle ou Mur", "Plancher Chauffant", "Intervention Rapide", "Toulouse et 25 km Autour",
      "Techniciens Équipés", "Compteur Qui Tourne ?", "Disponible 7j/7"],
     ["Fuite d'eau cachée ? Localisation précise sans casse : caméra thermique, gaz traceur.",
      "Rapport détaillé avec photos pour votre déclaration de sinistre dégât des eaux."] + COMMON_DESC[1:3],
     "Recherche-fuite", "Toulouse"),

    (C1, "Debouchage canalisation", "/debouchage-canalisation-toulouse/", 5.00, 4.00,
     ["debouchage canalisation toulouse", "deboucher canalisation toulouse", "canalisation bouchee toulouse", "debouchage toulouse",
      "debouchage evier toulouse", "evier bouche", "douche bouchee", "baignoire bouchee", "lavabo bouche", "hydrocurage toulouse",
      "plombier debouchage toulouse", "canalisation bouchee urgence", "colonne bouchee", "debouchage canalisation urgence"],
     ["Débouchage Canalisation", "Canalisation Bouchée ?", "Débouchage 24h/24 Toulouse", "Appelez le 09 70 70 71 29",
      "Évier, Douche, Baignoire", "Furet et Hydrocurage", "Devis Écrit Avant Travaux", "Intervention Rapide",
      "Inspection Caméra", "Nuit, Week-end et Fériés", "Toulouse et 25 km Autour", "Débouchage Propre",
      "Plombiers Qualifiés", "Refoulement ? On Intervient", "Paiement CB Accepté"],
     ["Évier, douche ou canalisation bouchée ? Débouchage au furet ou par hydrocurage, 24h/24.",
      "Bouchons répétés ? Inspection caméra pour trouver la cause et régler le problème."] + COMMON_DESC[1:3],
     "Debouchage", "Toulouse"),

    (C1, "WC bouche", "/debouchage-wc-toulouse/", 4.50, 3.60,
     ["wc bouche toulouse", "toilettes bouchees toulouse", "debouchage wc toulouse", "deboucher wc toulouse", "wc bouche urgence",
      "toilette bouchee", "wc bouche", "debouchage toilettes", "sanibroyeur bouche", "sanibroyeur en panne toulouse",
      "plombier wc bouche", "wc qui deborde"],
     ["WC Bouché à Toulouse ?", "Débouchage WC 24h/24", "Toilettes Bouchées ?", "Appelez le 09 70 70 71 29",
      "Intervention Propre", "Sanibroyeur Bloqué ?", "Devis Écrit Avant Travaux", "Intervention Rapide",
      "Nuit, Week-end et Fériés", "Toulouse et 25 km Autour", "Plombier en 30 min Environ", "Plombiers Qualifiés",
      "Paiement CB Accepté", "Débouchage Toilettes", "WC Qui Déborde ?"],
     ["WC bouché ou qui déborde ? Un plombier débouche vos toilettes rapidement, 24h/24.",
      "Débouchage et dépannage de sanibroyeur toutes marques, intervention propre."] + COMMON_DESC[1:3],
     "WC-bouche", "Toulouse"),

    (C1, "Chauffe-eau", "/depannage-chauffe-eau-toulouse/", 4.50, 3.60,
     ["chauffe eau en panne toulouse", "depannage chauffe eau toulouse", "reparation chauffe eau toulouse", "ballon d'eau chaude en panne",
      "plus d'eau chaude", "chauffe eau qui fuit", "cumulus en panne", "remplacement chauffe eau toulouse", "changer chauffe eau toulouse",
      "plombier chauffe eau toulouse", "chauffe eau ne chauffe plus", "groupe de securite chauffe eau fuit", "ballon eau chaude toulouse",
      "chauffe eau thermodynamique panne"],
     ["Chauffe-eau en Panne ?", "Dépannage Chauffe-eau 24h/24", "Plus d'Eau Chaude ?", "Appelez le 09 70 70 71 29",
      "Toutes Marques", "Atlantic, Thermor, Ariston", "Réparation ou Remplacement", "Devis Écrit Avant Travaux",
      "Ballon d'Eau Chaude", "Chauffe-eau Qui Fuit ?", "Toulouse et 25 km Autour", "Intervention Rapide",
      "Plombiers Qualifiés", "Résistance, Thermostat", "Paiement CB Accepté"],
     ["Plus d'eau chaude ? Dépannage et remplacement de chauffe-eau toutes marques, 24h/24.",
      "Résistance, thermostat, groupe de sécurité : diagnostic et devis clair sur place."] + COMMON_DESC[1:3],
     "Chauffe-eau", "Toulouse"),

    (C1, "Degat des eaux", "/degat-des-eaux-toulouse/", 4.50, 3.60,
     ["degat des eaux toulouse", "degat des eaux plombier", "degat des eaux urgence", "inondation appartement plombier",
      "rupture canalisation", "tuyau eclate", "tuyau perce", "flexible eclate", "degat des eaux voisin", "plombier degat des eaux toulouse",
      "infiltration d'eau plafond"],
     ["Dégât des Eaux à Toulouse ?", "Plombier Urgence 24h/24", "Fuite Stoppée Rapidement", "Appelez le 09 70 70 71 29",
      "Rapport pour l'Assurance", "Tuyau Éclaté ? On Intervient", "Devis Écrit Avant Travaux", "Nuit, Week-end et Fériés",
      "Recherche de Fuite", "Toulouse et 25 km Autour", "Intervention Rapide", "Plombiers Qualifiés",
      "Facture Détaillée", "Coupez l'Eau et Appelez", "Maison ou Appartement"],
     ["Dégât des eaux : on stoppe la fuite, on sécurise et on répare. Plombier 24h/24.",
      "Facture et rapport d'intervention détaillés pour votre déclaration à l'assurance."] + COMMON_DESC[1:3],
     "Degat-des-eaux", "Toulouse"),

    (C1, "Robinet", "/reparation-robinet-toulouse/", 2.50, 2.00,
     ["reparation robinet toulouse", "robinet qui fuit toulouse", "mitigeur qui fuit", "changer robinet toulouse",
      "remplacement mitigeur toulouse", "robinet casse", "robinet thermostatique douche fuit", "robinet d'arret fuit",
      "plombier robinet toulouse", "chasse d'eau qui fuit"],
     ["Robinet Qui Fuit ?", "Réparation Robinet Toulouse", "Mitigeur, Robinet, Chasse", "Appelez le 09 70 70 71 29",
      "Toutes Marques", "Grohe, Hansgrohe, Roca", "Devis Écrit Avant Travaux", "Remplacement Rapide",
      "Robinet Thermostatique", "Toulouse et 25 km Autour", "Intervention Rapide", "Plombiers Qualifiés",
      "Robinet d'Arrêt Grippé ?", "Chasse d'Eau Qui Fuit ?", "Paiement CB Accepté"],
     ["Robinet qui goutte, mitigeur cassé, chasse d'eau qui fuit : réparation ou remplacement.",
      "Joint, cartouche ou robinet complet : le plombier vous conseille et vous remet un devis."] + COMMON_DESC[1:3],
     "Robinet", "Toulouse"),
]

# Campagne communes : un groupe par commune principale + un groupe « autres communes »
MAIN_COMMUNES = ["Blagnac", "Colomiers", "Tournefeuille", "Balma", "L'Union", "Muret", "Cugnaux", "Ramonville-Saint-Agne",
                 "Saint-Orens-de-Gameville", "Plaisance-du-Touch", "Castanet-Tolosan", "Portet-sur-Garonne", "Villeneuve-Tolosane",
                 "Saint-Jean", "Aucamville", "Labège", "Fenouillet", "Launaguet", "Castelginest", "Cornebarrieu", "Pibrac"]


def kw_name(n):
    """Nom de commune tel que tapé dans Google (minuscules, sans apostrophe ni tiret)."""
    import unicodedata
    n = unicodedata.normalize("NFD", n.lower())
    n = "".join(ch for ch in n if unicodedata.category(ch) != "Mn")
    return n.replace("'", " ").replace("-", " ")


def short(n, limit):
    """Nom court pour les titres (ex. « Ramonville » pour Ramonville-Saint-Agne)."""
    alias = {"Ramonville-Saint-Agne": "Ramonville", "Saint-Orens-de-Gameville": "Saint-Orens", "Plaisance-du-Touch": "Plaisance-du-Touch",
             "Villeneuve-Tolosane": "Villeneuve-Tolosane", "Portet-sur-Garonne": "Portet", "Castanet-Tolosan": "Castanet"}
    n2 = alias.get(n, n)
    return n2 if len(n2) <= limit else n2.split("-")[0]


def commune_group(name):
    cp = next(c[1] for c in COMMUNES if c[0] == name)
    k = kw_name(name)
    s = short(name, 19)
    kws = ["plombier %s" % k, "plombier urgence %s" % k, "plombier a %s" % k, "plombier %s %s" % (k, cp), "plomberie %s" % k,
           "depannage plomberie %s" % k, "fuite d'eau %s" % k, "debouchage %s" % k, "wc bouche %s" % k, "chauffe eau %s" % k,
           "recherche de fuite %s" % k]
    def fit(limit, *options):
        return next(o for o in options if len(o) <= limit)
    heads = ["Plombier %s" % s, fit(30, "Plombier Urgence %s" % s, "Urgence Plombier %s" % s, "SOS Plombier %s" % s, "Urgence %s" % s),
             fit(30, "Plombier à %s 24h/24" % s, "Plombier %s 24h/24" % s, "%s 24h/24" % s),
             "Appelez le 09 70 70 71 29", "Devis Écrit Avant Travaux", "Intervention Rapide", "Fuite, Bouchon, Chauffe-eau",
             "Nuit, Week-end et Fériés", "Plombiers Qualifiés", "Facture pour l'Assurance", "SOS Plombier 24h/24 7j/7",
             fit(30, "Dépannage Plomberie %s" % s, "Dépannage %s" % s, "Plomberie %s" % s), "Toutes Marques", "Paiement CB Accepté",
             fit(30, "Plombier %s (%s)" % (s, cp), "Plombier %s" % cp)]
    descs = [fit(90, "Plombier à %s (%s) disponible 24h/24 : fuite d'eau, débouchage, chauffe-eau." % (s, cp),
                 "Plombier à %s (%s) 24h/24 : fuite d'eau, débouchage, chauffe-eau." % (s, cp),
                 "Plombier à %s 24h/24 : fuite d'eau, débouchage, chauffe-eau." % s),
             fit(90, "Un plombier se déplace à %s rapidement, de jour comme de nuit. Appelez !" % s,
                 "Plombier à %s rapidement, de jour comme de nuit. Appelez maintenant !" % s,
                 "Plombier à %s jour et nuit. Appelez maintenant !" % s),
             COMMON_DESC[1], COMMON_DESC[3]]
    return (C2, "Plombier %s" % name, "/", 4.00, 3.20, kws, heads, descs, "Plombier", short(name, 15).replace(" ", "-"))


def autres_group():
    others = [c for c in COMMUNES if c[0] not in MAIN_COMMUNES]
    kws = []
    for c in others:
        k = kw_name(c[0])
        kws += ["plombier %s" % k, "plombier urgence %s" % k]
    heads = ["Plombier Près de Chez Vous", "Plombier Urgence 24h/24", "Toulouse et 25 km Autour", "Appelez le 09 70 70 71 29",
             "Devis Écrit Avant Travaux", "Intervention Rapide", "Fuite, Bouchon, Chauffe-eau", "Nuit, Week-end et Fériés",
             "Plombiers Qualifiés", "Facture pour l'Assurance", "SOS Plombier Haute-Garonne", "Plombier dans Votre Commune",
             "Toutes Marques", "Paiement CB Accepté", "Dépannage Plomberie 7j/7"]
    descs = ["Plombier d'urgence dans votre commune, à 25 km autour de Toulouse, 24h/24 et 7j/7.",
             "Fuite d'eau, WC bouché, chauffe-eau en panne : un plombier se déplace rapidement.",
             COMMON_DESC[1], COMMON_DESC[3]]
    return (C2, "Plombier Autres communes", "/zones-intervention/", 3.50, 2.80, kws, heads, descs, "Plombier", "Haute-Garonne")


NEGATIVES = [
    # Emploi / formation
    "emploi", "offre d'emploi", "recrutement", "recrute", "embauche", "salaire", "salaires", "cdi", "cdd", "interim", "interimaire",
    "alternance", "apprenti", "apprentissage", "stage", "stagiaire", "formation", "cap plombier", "bac pro", "devenir plombier",
    "ecole", "cours", "metier", "fiche metier", "candidature", "cv", "pole emploi", "france travail", "indeed", "auto entrepreneur",
    # Bricolage / do it yourself
    "tuto", "tutoriel", "comment", "video", "youtube", "soi meme", "moi meme", "diy", "astuce", "astuces", "bicarbonate", "vinaigre",
    "schema", "notice", "mode d'emploi", "pdf", "norme", "dtu",
    # Achat de matériel / magasins
    "leroy merlin", "castorama", "brico depot", "bricomarche", "mr bricolage", "manomano", "amazon", "cedeo", "point p", "brossette",
    "outil", "outils", "outillage", "piece detachee", "pieces detachees", "acheter", "achat", "vente", "occasion", "location",
    "pvc prix", "tube per", "pince", "cle a molette", "furet manuel", "ventouse", "destop", "deboucheur",
    # Hors sujet / autres métiers
    "mario", "jeu", "jeux", "film", "dessin", "coloriage", "definition", "wikipedia", "synonyme", "humour", "blague",
    "electricien", "serrurier", "climatisation", "pompe a chaleur", "piscine", "vidange fosse septique", "fosse septique",
    "ramonage", "chaudiere gaz entretien", "panneau solaire",
    # Gratuit / plaintes
    "arnaque", "plainte", "litige", "avis negatif", "escroquerie",
    # Autres villes (hors zone)
    "paris", "lyon", "marseille", "bordeaux", "montpellier", "nice", "nantes", "lille", "strasbourg", "albi", "montauban",
    "castres", "carcassonne", "pamiers", "foix", "auch", "tarbes", "pau", "agen", "rodez", "cahors", "saint gaudens", "revel",
]

SITELINKS = [
    ("Fuite d'eau", "/fuite-d-eau-toulouse/", "Fuite stoppée rapidement", "Réparation 24h/24"),
    ("Recherche de fuite", "/recherche-de-fuite-toulouse/", "Détection sans casse", "Rapport pour l'assurance"),
    ("Débouchage canalisation", "/debouchage-canalisation-toulouse/", "Évier, douche, baignoire", "Furet et hydrocurage"),
    ("WC bouché", "/debouchage-wc-toulouse/", "Débouchage de toilettes", "Sanibroyeur toutes marques"),
    ("Chauffe-eau en panne", "/depannage-chauffe-eau-toulouse/", "Plus d'eau chaude ?", "Toutes marques"),
    ("Dégât des eaux", "/degat-des-eaux-toulouse/", "Intervention d'urgence", "Facture pour l'assurance"),
    ("Devis gratuit en ligne", "/#contact", "Réponse rapide", "Sans engagement"),
    ("Zones d'intervention", "/zones-intervention/", "Toulouse et 25 km autour", "Tous les quartiers"),
]
CALLOUTS = ["Disponible 24h/24 7j/7", "Devis Avant Travaux", "Intervention Rapide", "Nuit, Week-end, Fériés", "Toutes Marques",
            "Paiement CB Accepté", "Facture pour l'Assurance", "Plombiers Qualifiés", "Détection Sans Casse", "Toulouse + 25 km"]
SNIPPET_SERVICES = ["Fuite d'eau", "Recherche de fuite", "Débouchage", "WC bouché", "Chauffe-eau", "Dégât des eaux", "Robinetterie"]


def check(label, items, limit):
    for t in items:
        if len(t) > limit:
            raise SystemExit("Trop long (%d > %d) – %s : %r" % (len(t), limit, label, t))


def main():
    os.makedirs(OUT, exist_ok=True)
    groups = AD_GROUPS + [commune_group(n) for n in MAIN_COMMUNES] + [autres_group()]

    for g in groups:
        camp, name, url, bex, bph, kws, heads, descs, p1, p2 = g
        check(name + " titres", heads, 30)
        check(name + " descriptions", descs, 90)
        check(name + " chemins", [p1, p2], 15)
        assert len(heads) == 15, (name, len(heads))
        assert len(descs) == 4, (name, len(descs))
        assert len(set(h.lower() for h in heads)) == 15, ("titres en double", name)
    check("liens annexes", [s[0] for s in SITELINKS], 25)
    check("liens annexes desc", [d for s in SITELINKS for d in s[2:]], 35)
    check("accroches", CALLOUTS, 25)
    check("extrait", SNIPPET_SERVICES, 25)

    def w(name, header, rows):
        with open(os.path.join(OUT, name), "w", encoding="utf-8-sig", newline="") as f:
            wr = csv.writer(f)
            wr.writerow(header)
            wr.writerows(rows)

    # 1. Campagnes
    w("1-campagnes.csv", ["Campaign", "Campaign Type", "Networks", "Budget", "Budget type", "Bid Strategy Type", "Languages", "Location", "Campaign Status"],
      [[C1, "Search", "Google search", "50.00", "Daily", "Manual CPC", "fr", "Toulouse (rayon 25 km)", "Paused"],
       [C2, "Search", "Google search", "20.00", "Daily", "Manual CPC", "fr", "Toulouse (rayon 25 km)", "Paused"]])

    # 2. Mots-clés (exact + expression) avec enchères
    rows = []
    n_kw = 0
    for camp, name, url, bex, bph, kws, *_ in groups:
        for k in kws:
            rows.append([camp, name, k, "Exact", "%.2f" % bex, BASE + url])
            n_kw += 1
            if bph:
                rows.append([camp, name, k, "Phrase", "%.2f" % bph, BASE + url])
                n_kw += 1
    w("2-mots-cles.csv", ["Campaign", "Ad Group", "Keyword", "Criterion Type", "Max CPC", "Final URL"], rows)

    # 3. Groupes d'annonces
    w("3-groupes-annonces.csv", ["Campaign", "Ad Group", "Max CPC", "Ad Group Status"],
      [[g[0], g[1], "%.2f" % g[3], "Enabled"] for g in groups])

    # 4. Annonces responsives (RSA)
    head = ["Campaign", "Ad Group", "Ad type"] + ["Headline %d" % i for i in range(1, 16)] + ["Headline 1 position"] + \
           ["Description %d" % i for i in range(1, 5)] + ["Path 1", "Path 2", "Final URL"]
    rows = []
    for camp, name, url, bex, bph, kws, heads, descs, p1, p2 in groups:
        rows.append([camp, name, "Responsive search ad"] + heads + ["1"] + descs + [p1, p2, BASE + url])
    w("4-annonces-rsa.csv", head, rows)

    # 5. Mots-clés négatifs (niveau campagne, en expression)
    rows = [[c, n, "Campaign Negative Phrase"] for c in (C1, C2) for n in NEGATIVES]
    w("5-mots-cles-negatifs.csv", ["Campaign", "Keyword", "Criterion Type"], rows)

    # 6. Composants (extensions)
    rows = [["Sitelink", c, s[0], s[2], s[3], BASE + s[1]] for c in (C1, C2) for s in SITELINKS]
    rows += [["Callout", c, t, "", "", ""] for c in (C1, C2) for t in CALLOUTS]
    rows += [["Structured snippet (Services)", c, "; ".join(SNIPPET_SERVICES), "", "", ""] for c in (C1, C2)]
    rows += [["Call", c, SITE["phone_display"], "Pays : France", "Rapports d'appels : activés", "Conversion : appel >= 60 s"] for c in (C1, C2)]
    w("6-composants-extensions.csv", ["Type", "Campaign", "Texte", "Ligne 1", "Ligne 2", "URL / Option"], rows)

    write_xlsx(groups, n_kw)
    print("%d groupes d'annonces, %d mots-clés, %d négatifs : OK (limites de caractères vérifiées)" % (len(groups), n_kw, len(NEGATIVES)))


def write_xlsx(groups, n_kw):
    """Classeur Excel lisible : un onglet par élément de la campagne."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_fill = PatternFill("solid", fgColor="1A2E2A")
    head_font = Font(bold=True, color="FFFFFF")
    blue = PatternFill("solid", fgColor="EFF6FF")
    wb = Workbook()

    def sheet(title, header, rows, widths=None, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, wdt in enumerate(widths or [], 1):
            ws.column_dimensions[get_column_letter(i)].width = wdt
        return ws

    # Résumé
    ws = sheet("Résumé", ["Élément", "Valeur"], [], [38, 90], first=True)
    budgets = {C1: 50, C2: 20}
    rows = [
        ("Compte", "Artisans France Plomberie – " + SITE["base_url"]),
        ("Stratégie d'enchères", "CPC manuel (sans requête large)"),
        ("Réseaux", "Réseau de Recherche Google uniquement (décocher Partenaires et Display)"),
        ("Zone", "Toulouse + rayon de 25 km – option « Présence » uniquement"),
        ("Langue", "Français"),
        ("Calendrier", "24h/24, 7j/7"),
        ("Ajustements appareils", "Mobile +20 %, Ordinateur –20 %, Tablette –30 %"),
        ("Ajustements horaires", "19h–8h +15 %, samedi et dimanche +10 %"),
        ("Campagne 1", "%s – budget %d €/jour – %d groupes" % (C1, budgets[C1], sum(1 for g in groups if g[0] == C1))),
        ("Campagne 2", "%s – budget %d €/jour – %d groupes" % (C2, budgets[C2], sum(1 for g in groups if g[0] == C2))),
        ("Mots-clés", "%d (exact + expression)" % n_kw),
        ("Mots-clés négatifs", "%d par campagne" % len(NEGATIVES)),
        ("Conversions à suivre", "Appel depuis l'annonce ≥ 60 s · Clic sur le numéro du site · Formulaire de devis envoyé"),
        ("Téléphone (extension d'appel)", SITE["phone_display"]),
        ("Important", "Les enchères sont des points de départ : à vérifier dans le Planificateur de mots clés puis à ajuster après 7 à 14 jours."),
    ]
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    # Groupes d'annonces
    sheet("Groupes d'annonces", ["Campagne", "Groupe d'annonces", "Enchère exacte (€)", "Enchère expression (€)", "Nb mots-clés", "Page de destination"],
          [[g[0], g[1], g[3], g[4] or "–", len(g[5]), BASE + g[2]] for g in groups], [36, 34, 18, 20, 13, 62])

    # Mots-clés
    rows = []
    for camp, name, url, bex, bph, kws, *_ in groups:
        for k in kws:
            rows.append([camp, name, k, "Exact", "[%s]" % k, bex, BASE + url])
            if bph:
                rows.append([camp, name, k, "Expression", '"%s"' % k, bph, BASE + url])
    ws = sheet("Mots-clés", ["Campagne", "Groupe d'annonces", "Mot-clé", "Correspondance", "Format Google Ads", "CPC max (€)", "URL finale"], rows, [34, 30, 44, 15, 46, 12, 60])
    for c in ws["F"][1:]:
        c.number_format = '0.00 "€"'

    # Annonces : une ligne par élément avec nombre de caractères
    rows = []
    for camp, name, url, bex, bph, kws, heads, descs, p1, p2 in groups:
        for i, h in enumerate(heads, 1):
            rows.append([camp, name, "Titre %d%s" % (i, " (épinglé position 1)" if i == 1 else ""), h, len(h), 30])
        for i, d in enumerate(descs, 1):
            rows.append([camp, name, "Description %d" % i, d, len(d), 90])
        rows.append([camp, name, "Chemin 1", p1, len(p1), 15])
        rows.append([camp, name, "Chemin 2", p2, len(p2), 15])
        rows.append([camp, name, "URL finale", BASE + url, "", ""])
    ws = sheet("Annonces RSA", ["Campagne", "Groupe d'annonces", "Élément", "Texte", "Caractères", "Limite"], rows, [34, 30, 26, 92, 11, 8])
    for row in ws.iter_rows(min_row=2):
        if str(row[2].value).startswith("Description"):
            for c in row:
                c.fill = blue

    # Négatifs
    sheet("Mots-clés négatifs", ["Campagne", "Mot-clé négatif", "Correspondance", "Format Google Ads"],
          [[c, n, "Expression", '"%s"' % n] for c in (C1, C2) for n in NEGATIVES], [36, 34, 16, 36])

    # Extensions
    rows = [["Lien annexe", s[0], s[2], s[3], BASE + s[1]] for s in SITELINKS]
    rows += [["Accroche", t, "", "", ""] for t in CALLOUTS]
    rows += [["Extrait de site (Services)", v, "", "", ""] for v in SNIPPET_SERVICES]
    rows += [["Appel", SITE["phone_display"], "Rapports d'appels activés", "Conversion : appel ≥ 60 s", ""]]
    sheet("Extensions", ["Type", "Texte", "Description 1", "Description 2", "URL"], rows, [26, 30, 30, 30, 62])

    # Format d'import Google Ads Editor (copier-coller)
    head = ["Campaign", "Ad Group", "Ad type"] + ["Headline %d" % i for i in range(1, 16)] + \
           ["Description %d" % i for i in range(1, 5)] + ["Path 1", "Path 2", "Final URL"]
    sheet("Import Editor - RSA", head, [[g[0], g[1], "Responsive search ad"] + g[6] + g[7] + [g[8], g[9], BASE + g[2]] for g in groups])

    wb.save(os.path.join(OUT, "campagne-google-ads-plombier-toulouse.xlsx"))


if __name__ == "__main__":
    main()
